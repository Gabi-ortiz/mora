"""Normaliza y clasifica un extracto Banco Macro (CSV) y genera una planilla de trabajo.

Uso: python procesar_extracto.py extracto.csv salida.xlsx
"""
import csv
import re
import sys
from collections import defaultdict
from datetime import datetime

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

from reglas_macro import BANCO, REGLAS

CUIT_RE = re.compile(r"\b(20|23|24|27|30|33|34)(\d{8})(\d)\b")
NUM_FMT = '#,##0.00;[Red]-#,##0.00'
HEADER_FILL = PatternFill("solid", fgColor="1F4E78")


PREGUNTAS = [
    ("CHEXA SA (30681026386)", "¿Es la titular de esta cuenta o una empresa del grupo? Define si 2018/3808/4254/4267 son movimientos internos."),
    ("Datanet M/T y BTOB (2018, 2026)", "¿Hacia qué cuentas van las transferencias 'ABONADO M/T' y 'BTOB'? ¿Cómo se registran en FBS?"),
    ("Pagos a proveedores (2027)", "Los débitos 'TRANSF DATANET PROVEEDORES' son lotes (hasta $1.000M). ¿FBS tiene el detalle por proveedor de cada lote?"),
    ("IMP. AFIP $1.770 (3696, 4196)", "63 débitos de $1.770 c/u. ¿Qué son (VEP, comisión, tasa)? ¿Se registran uno por uno o agrupados?"),
    ("Número de Operación (4333, 4334)", "Créditos 'NNN - Numero de Operacion'. ¿Son cobranzas por el convenio de recaudación 61068? ¿Cómo se identifica al cliente?"),
    ("DL1 (366)", "Créditos '<CUIT> <NOMBRE> DL1' por montos grandes. ¿Qué operatoria es?"),
    ("PP- N/C Deuda Pública (1952)", "¿Qué corresponde a 'PP- N/C DEUD. PUBLICA-P.PREVIO'?"),
    ("APC Aportes (4218, 4220)", "Transferencias 'BURZIO A - APC APORTES DE C'. ¿Aportes de capital? ¿Cómo se contabilizan?"),
    ("Banco del Sol / VAZCAR (4061, 4066)", "¿Son préstamos, cobranzas o transferencias propias?"),
    ("Plan de ahorro Chevrolet (4196, 1445)", "¿Los débitos 'PLAN AHORRO CHEVROLET' y créditos 'TRPRO CHEVROLET' se registran contra la cuenta de la terminal en FBS?"),
    ("Cobranzas clientes", "¿En FBS la cobranza se registra por cliente con CUIT? Permitiría el cruce automático por CUIT + importe + fecha."),
    ("Impuestos y comisiones", "¿Ley 25.413, SIRCREB y comisiones se cargan en FBS por movimiento, por día o por mes?"),
    ("Cheques", "En julio no hay cheques en el extracto. ¿Esta cuenta emite/recibe cheques o eCheq? ¿Dónde se controlan los pendientes?"),
    ("Datos FBS", "¿Qué exportación de FBS podemos usar (mayor de la cuenta banco Macro) y en qué formato?"),
]


def num(s):
    return float(s.replace(".", "").replace(",", ".")) if s.strip() else 0.0


def cuit_valido(c):
    pesos = [5, 4, 3, 2, 7, 6, 5, 4, 3, 2]
    r = 11 - sum(int(d) * p for d, p in zip(c[:10], pesos)) % 11
    return int(c[10]) == {11: 0, 10: 9}.get(r, r)


def extraer_cuit(concepto):
    for m in CUIT_RE.finditer(concepto):
        c = "".join(m.groups())
        if cuit_valido(c):
            return c
    return ""


def extraer_nombre(concepto):
    for pat in (r"ING TRANSF:(.+?)-\d{11}", r"^TRANSF (.+?) \d{11}", r"TEF DATANET \w+ (.+?)\s+\d{11}",
                r"^CCERR (.+?) \d{11}", r"^TRMIN (.+?) \d{11}", r"^TRPRO (.+?) \d{11}",
                r"^\d{11} (.+?) DL1", r"Transf\. \w+ (.+?)\s+\d{11}"):
        m = re.search(pat, concepto)
        if m:
            return re.sub(r"\s+", " ", m.group(1)).strip(" /,")
    return ""


def clasificar(causal, concepto):
    for cod, patron, cat, sub, fbs in REGLAS:
        if cod == causal and (patron is None or re.search(patron, concepto)):
            return cat, sub, fbs
    return "A IDENTIFICAR", "Código causal nuevo", ""


def leer(path):
    with open(path, encoding="utf-8-sig", newline="") as f:
        rows = list(csv.DictReader(f, delimiter=";"))
    movs = []
    for r in reversed(rows):  # el banco lo entrega del más nuevo al más viejo
        deb, cred = num(r["Debito"]), num(r["Credito"])
        cat, sub, fbs = clasificar(r["Codigo Causal"], r["Concepto"])
        movs.append({
            "fecha": datetime.strptime(r["Fecha"], "%d/%m/%Y").date(),
            "ref": r["Referencia"], "causal": r["Codigo Causal"], "concepto": r["Concepto"].strip(),
            "debito": deb, "credito": cred, "importe": cred - deb, "saldo": num(r["Saldo"]),
            "cat": cat, "sub": sub, "fbs": fbs,
            "cuit": extraer_cuit(r["Concepto"]), "nombre": extraer_nombre(r["Concepto"]),
        })
    return movs


def hoja(wb, titulo, headers, filas, anchos, formatos=None):
    ws = wb.create_sheet(titulo)
    ws.append(headers)
    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = HEADER_FILL
        c.alignment = Alignment(wrap_text=True, vertical="center")
    for f in filas:
        ws.append(f)
    for i, w in enumerate(anchos, 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    for col, fmt in (formatos or {}).items():
        for cell in ws[col][1:]:
            cell.number_format = fmt
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    return ws


def generar(movs, salida):
    wb = Workbook()
    wb.remove(wb.active)

    # Resumen y control de saldos
    ini = movs[0]["saldo"] - movs[0]["importe"]
    tot_c = sum(m["credito"] for m in movs)
    tot_d = sum(m["debito"] for m in movs)
    fin = movs[-1]["saldo"]
    quiebres = sum(1 for a, b in zip(movs, movs[1:]) if abs(a["saldo"] + b["importe"] - b["saldo"]) > 0.005)
    por_cat = defaultdict(lambda: [0, 0.0, 0.0])
    for m in movs:
        x = por_cat[m["cat"]]
        x[0] += 1; x[1] += m["credito"]; x[2] += m["debito"]
    ws = hoja(wb, "Resumen", ["Concepto", "Valor"], [
        ["Banco", BANCO],
        ["Período", f"{movs[0]['fecha']:%d/%m/%Y} al {movs[-1]['fecha']:%d/%m/%Y}"],
        ["Cantidad de movimientos", len(movs)],
        ["Saldo inicial (calculado)", ini],
        ["(+) Créditos", tot_c],
        ["(-) Débitos", tot_d],
        ["Saldo final según extracto", fin],
        ["Diferencia control (debe ser 0)", round(ini + tot_c - tot_d - fin, 2)],
        ["Quiebres de saldo entre filas", quiebres],
        [], ["Categoría", "Cant.", "Créditos", "Débitos", "Neto"],
    ] + [[k, v[0], v[1], v[2], v[1] - v[2]] for k, v in sorted(por_cat.items(), key=lambda kv: -(kv[1][1] + kv[1][2]))],
        [42, 20, 20, 20, 20])
    for row in ws.iter_rows(min_row=2):
        for c in row[1:]:
            if isinstance(c.value, float):
                c.number_format = NUM_FMT
    ws["A12"].font = ws["B12"].font = Font(bold=True)

    # Catálogo de conceptos (una fila por regla, con estadísticas del mes)
    stats = defaultdict(lambda: [0, 0.0, 0.0, []])
    for m in movs:
        k = (m["causal"], m["sub"])
        s = stats[k]
        s[0] += 1; s[1] += m["debito"]; s[2] += m["credito"]
        if len(s[3]) < 3 and m["concepto"] not in s[3]:
            s[3].append(m["concepto"])
    filas = []
    for cod, patron, cat, sub, fbs in REGLAS:
        s = stats.get((cod, sub))
        if not s:
            continue
        filas.append([BANCO, int(cod), patron or "(todos)", cat, sub, s[0], s[1], s[2], " | ".join(s[3]),
                      fbs, "A validar", ""])
    for (cod, sub), s in stats.items():
        if sub == "Código causal nuevo":
            filas.append([BANCO, int(cod), "(todos)", "A IDENTIFICAR", sub, s[0], s[1], s[2], " | ".join(s[3]), "", "A validar", ""])
    hoja(wb, "Catalogo_Conceptos", ["Banco", "Cód. causal", "Patrón en concepto", "Categoría", "Descripción propuesta",
                                    "Cant. mes", "Débitos mes", "Créditos mes", "Ejemplos del extracto",
                                    "Registro esperado en FBS", "Estado validación", "Comentario administración"],
         filas, [8, 10, 16, 34, 42, 9, 18, 18, 70, 34, 14, 40], {"G": NUM_FMT, "H": NUM_FMT})

    # Movimientos normalizados
    hoja(wb, "Movimientos", ["Fecha", "Referencia", "Cód. causal", "Concepto", "Débito", "Crédito", "Importe (+/-)",
                             "Saldo", "Categoría", "Descripción", "CUIT contraparte", "Nombre contraparte",
                             "Estado conciliación", "Comprobante FBS", "Observaciones"],
         [[m["fecha"], m["ref"], int(m["causal"]), m["concepto"], m["debito"] or None, m["credito"] or None,
           m["importe"], m["saldo"], m["cat"], m["sub"], m["cuit"], m["nombre"], "Pendiente", "", ""] for m in movs],
         [11, 13, 9, 48, 16, 16, 16, 18, 30, 34, 14, 28, 16, 18, 30],
         {"A": "dd/mm/yyyy", "E": NUM_FMT, "F": NUM_FMT, "G": NUM_FMT, "H": NUM_FMT})

    hoja(wb, "Preguntas_Admin", ["#", "Tema", "Pregunta", "Respuesta administración"],
         [[i, t, q, ""] for i, (t, q) in enumerate(PREGUNTAS, 1)], [4, 26, 100, 50])

    wb.save(salida)


if __name__ == "__main__":
    generar(leer(sys.argv[1]), sys.argv[2])
