"""Conciliación automática Banco Macro vs. mayor FBS (cuentas E + O).

Uso:
    python conciliar.py EXTRACTO MAYOR_E.pdf MAYOR_O.pdf SALIDA.xlsx [CONCILIACION_ANTERIOR.xlsx]

CONCILIACION_ANTERIOR es la planilla manual de conciliación (hoja "Macro"); de ella se toman las partidas
pendientes con fecha anterior al período (arrastre). Más adelante el arrastre será la hoja de pendientes
que genera este mismo script.

Criterios de cruce, en orden de prioridad (importe y sentido siempre deben coincidir):
  1. CUIT / DNI       - CUIT del extracto contra CUIT (TR30 50004045 5) o DNI (SX 32458203) del comentario FBS
  2. Referencia       - número del extracto (ej. cheque 64274954) presente en el comprobante FBS
  3. Nombre           - apellido de la contraparte en ambos lados
  4. Fecha + importe  - mismo día y mismo importe, candidato único
  5. Importe          - mismo importe a pocos días de distancia (queda como "Sugerido", revisar)
Además: agrupados del mismo CUIT (varios recibos = una transferencia) y el asiento mensual de gastos bancarios.
"""
import re
import sys
from collections import defaultdict
from datetime import date, datetime, timedelta
from itertools import combinations

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

from leer_mayor_fbs import leer_mayor
from procesar_extracto import NUM_FMT, HEADER_FILL, extraer_cuit, hoja, leer

VENTANA_DIAS = 10
VENTANA_SUGERIDO = 10
VENTANA_ID = 31        # CUIT/DNI + importe es un criterio fuerte: se admite todo el mes
TOLERANCIA = 0.02      # redondeos de centavos entre banco y FBS
SECTORES = [
    ("S1", "Depósitos no registrados en contabilidad", "+"),
    ("S2", "Débitos no registrados en contabilidad", "-"),
    ("S3", "Depósitos no acreditados en banco", "-"),
    ("S4", "Cheques / pagos no debitados en banco", "+"),
]
HOJAS = {"S1": "1 Dep no registrados", "S2": "2 Deb no registrados", "S3": "3 Dep no acreditados",
         "S4": "4 Pagos no debitados"}
PALABRAS_VACIAS = {"SA", "SRL", "SAS", "S", "A", "DE", "LA", "EL", "Y", "BANCO", "MACRO", "TRANSF", "VAR", "VARIOS",
                   "CUO", "CUOTAS", "FAC", "FACTURAS", "TEF", "DATANET", "PR", "ING", "CIRC", "CERRADO", "CCERR"}


# ---------------------------------------------------------------- partidas

def _tokens(texto):
    return {t for t in re.findall(r"[A-ZÑ]{3,}", (texto or "").upper()) if t not in PALABRAS_VACIAS}


def partida_banco(m, origen="Mes"):
    cuit = m.get("cuit") or extraer_cuit(m["concepto"])
    return {"lado": "BANCO", "origen": origen, "fecha": m["fecha"], "importe": round(m["importe"], 2),
            "texto": m["concepto"], "ref": str(m.get("ref") or ""), "cuit": cuit, "dni": cuit[2:10] if cuit else "",
            "nombre": _tokens(m.get("nombre") or m["concepto"]), "cat": m.get("cat", ""), "causal": m.get("causal", ""),
            "match": None, "metodo": ""}


def partida_fbs(a, origen="Mes"):
    com = a["comentario"]
    cuit = dni = ""
    m = re.search(r"TR ?(\d{2})[ -]?(\d{8})[ -]?(\d)\b", com + " " + a.get("referencia", ""))
    if m:
        cuit = "".join(m.groups()); dni = m.group(2)
    else:
        m = re.search(r"\b(?:SX|TR) ?(\d{7,8})\b", com)
        if m:
            dni = m.group(1).zfill(8)
    partes = [p.strip() for p in com.split("/")]
    nombre = _tokens(partes[-1]) if len(partes) > 1 else set()
    return {"lado": "FBS", "origen": origen, "fecha": a["fecha"], "importe": round(a["debe"] - a["haber"], 2),
            "texto": com, "ref": a.get("referencia", ""), "asiento": a.get("asiento", ""), "cuit": cuit, "dni": dni,
            "nombre": nombre, "cat": "", "causal": "", "match": None, "metodo": ""}


def movimientos_fbs(asientos_e, asientos_o):
    """Une E y O descartando los pases entre ambas cuentas (mismo asiento e importe con signo opuesto)."""
    neto = defaultdict(float)
    filas = {}
    for a in asientos_e + asientos_o:
        k = (a["asiento"], a["comentario"], round(abs(a["debe"] - a["haber"]), 2))
        neto[k] += a["debe"] - a["haber"]
        filas.setdefault(k, a)
    out = []
    for k, v in neto.items():
        if abs(v) > 0.005:
            a = dict(filas[k]); a["debe"], a["haber"] = max(v, 0.0), max(-v, 0.0)
            out.append(a)
    return out


def _bloques_manual(ws):
    """Ubica los 4 sectores en la hoja de conciliación manual: (sector, columna inicial, fila desde, fila hasta)."""
    fila_s1 = fila_s3 = None
    for r in range(1, ws.max_row + 1):
        v = str(ws.cell(r, 2).value or "")
        if v.startswith("Depositos No Registrados"):
            fila_s1 = r
        elif v.startswith("Depositos No Acreditados"):
            fila_s3 = r
    fin_s3 = next(r for r in range(fila_s3, ws.max_row + 1) if str(ws.cell(r, 2).value or "").startswith("Saldo Contab"))
    return [("S1", 1, fila_s1 + 1, fila_s3 - 3), ("S2", 5, fila_s1 + 1, fila_s3 - 3),
            ("S3", 1, fila_s3 + 1, fin_s3 - 2), ("S4", 5, fila_s3 + 1, fin_s3 - 2)]


def leer_totales_manual(path):
    """Totales de la conciliación manual (para comparar durante la prueba piloto)."""
    ws = load_workbook(path, data_only=True)["Macro"]
    tot = {"E": ws["G3"].value, "O": ws["G4"].value}
    for sec, c0, r0, r1 in _bloques_manual(ws):
        tot[sec] = sum(v for r in range(r0, r1) if isinstance(v := ws.cell(r, c0 + 2).value, (int, float))
                       and (ws.cell(r, c0).value or ws.cell(r, c0 + 1).value))   # excluye la fila de total
    for r in range(1, ws.max_row + 1):
        if str(ws.cell(r, 2).value or "").startswith("Saldo Extracto"):
            tot["banco"] = ws.cell(r, 5).value
    return tot


def leer_arrastre(path, desde):
    """Partidas pendientes de la conciliación manual (hoja Macro) con fecha anterior al período."""
    ws = load_workbook(path, data_only=True)["Macro"]
    partidas = []
    for sec, c0, r0, r1 in _bloques_manual(ws):
        for r in range(r0, r1):
            f, texto = ws.cell(r, c0).value, ws.cell(r, c0 + 1).value
            imps = [ws.cell(r, c).value for c in (c0 + 2, c0 + 3)]
            imp = next((x for x in imps if isinstance(x, (int, float))), None)
            if imp is None or not (texto or f) or (isinstance(f, datetime) and f.date() >= desde):
                continue
            f = f.date() if isinstance(f, datetime) else desde - timedelta(days=1)   # partidas sin fecha
            m = {"fecha": f, "concepto": str(texto or ""), "importe": 0.0, "ref": "", "nombre": ""}
            if sec in ("S1", "S2"):     # vienen del banco
                m["importe"] = imp if sec == "S1" else -imp
                p = partida_banco(m, "Arrastre")
            else:                       # vienen de FBS
                p = partida_fbs({"fecha": f, "comentario": m["concepto"], "referencia": "",
                                 "debe": imp if sec == "S3" else 0.0, "haber": imp if sec == "S4" else 0.0}, "Arrastre")
            partidas.append(p)
    return partidas


# ---------------------------------------------------------------- cruce

def _dias(a, b):
    return abs((a["fecha"] - b["fecha"]).days)


def _unir(banco, fbs, metodo):
    for b in banco:
        b["match"], b["metodo"] = fbs, metodo
    for f in fbs:
        f["match"], f["metodo"] = banco, metodo


def _pasada(banco, fbs, criterio, metodo, ventana=VENTANA_DIAS, unico=False):
    libres_f = defaultdict(list)
    for f in fbs:
        if f["match"] is None:
            libres_f[round(f["importe"])].append(f)
    for b in sorted(banco, key=lambda x: x["fecha"]):
        if b["match"] is not None:
            continue
        cands = [f for k in (round(b["importe"]) - 1, round(b["importe"]), round(b["importe"]) + 1)
                 for f in libres_f[k] if f["match"] is None and abs(f["importe"] - b["importe"]) <= TOLERANCIA
                 and _dias(b, f) <= ventana and criterio(b, f)]
        if not cands or (unico and len(cands) > 1):
            continue
        f = min(cands, key=lambda f: (_dias(b, f), abs(f["importe"] - b["importe"])))
        _unir([b], [f], metodo)


def _agrupados(banco, fbs):
    """Una transferencia del banco = varios recibos FBS del mismo CUIT/DNI o apellido (o al revés)."""
    for uno, varios, lado_uno in ((banco, fbs, "BANCO"), (fbs, banco, "FBS")):
        for u in uno:
            if u["match"] is not None or not (u["dni"] or u["nombre"]):
                continue
            cands = [v for v in varios if v["match"] is None and _dias(u, v) <= VENTANA_DIAS
                     and (v["importe"] > 0) == (u["importe"] > 0)
                     and ((u["dni"] and v["dni"] == u["dni"]) or (u["nombre"] & v["nombre"]))][:12]
            for n in range(2, min(len(cands), 6) + 1):
                combo = next((c for c in combinations(cands, n)
                              if abs(sum(x["importe"] for x in c) - u["importe"]) <= TOLERANCIA), None)
                if combo:
                    if lado_uno == "BANCO":
                        _unir([u], list(combo), "Agrupado CUIT/DNI/Nombre")
                    else:
                        _unir(list(combo), [u], "Agrupado CUIT/DNI/Nombre")
                    break


def _combinaciones(banco, fbs, dias=7, max_items=3):
    """Un movimiento = combinación de 2-3 del otro lado a pocos días (ej. venta de cheques neta de gastos)."""
    for uno, varios, lado_uno in ((banco, fbs, "BANCO"), (fbs, banco, "FBS")):
        for u in sorted(uno, key=lambda x: -abs(x["importe"])):
            if u["match"] is not None or u["origen"] != "Mes":
                continue
            cands = sorted((v for v in varios if v["match"] is None and v["origen"] == "Mes" and _dias(u, v) <= dias),
                           key=lambda v: _dias(u, v))[:40]
            for n in range(2, max_items + 1):
                combo = next((c for c in combinations(cands, n)
                              if abs(sum(x["importe"] for x in c) - u["importe"]) <= TOLERANCIA), None)
                if combo:
                    if lado_uno == "BANCO":
                        _unir([u], list(combo), "Combinación (revisar)")
                    else:
                        _unir(list(combo), [u], "Combinación (revisar)")
                    break


def _compensaciones_fbs(fbs, misma_clave):
    """Registros FBS que se anulan entre sí el mismo día (reclasificaciones): no pasan por el banco.
    misma_clave=True solo compensa registros del mismo lote/comprobante (se corre antes del cruce con el banco)."""
    libres = defaultdict(list)
    for f in fbs:
        if f["match"] is None and f["origen"] == "Mes":
            libres[(f["fecha"], _clave_lote(f) if misma_clave else "")].append(f)
    for dia in libres.values():
        for u in sorted(dia, key=lambda x: -abs(x["importe"])):
            if u["match"] is not None:
                continue
            cands = [v for v in dia if v["match"] is None and (v["importe"] > 0) != (u["importe"] > 0)][:25]
            for n in range(1, 5):
                combo = next((c for c in combinations(cands, n)
                              if abs(sum(x["importe"] for x in c) + u["importe"]) <= TOLERANCIA), None)
                if combo:
                    for x in (u, *combo):
                        x["match"], x["metodo"] = [], "Compensa dentro de FBS"
                    break


def _reversiones_fbs(fbs):
    """Comprobante FBS registrado y luego revertido (mismo texto, importe opuesto, cualquier día del mes)."""
    libres = defaultdict(list)
    for f in fbs:
        if f["match"] is None and f["origen"] == "Mes":
            libres[(f["texto"], abs(f["importe"]))].append(f)
    for grupo in libres.values():
        for a, b in zip([f for f in grupo if f["importe"] > 0], [f for f in grupo if f["importe"] < 0]):
            a["match"], b["match"] = [], []
            a["metodo"] = b["metodo"] = "Compensa dentro de FBS"


def _clave_lote(f):
    m = re.search(r"Nº\s*(.+)$", f["texto"])
    return m.group(1).strip() if m else re.split(r"[\s/-]", f["texto"])[0]


def _lotes(banco, fbs, dias=3):
    """N movimientos del banco (mismo día y causal) = M registros FBS (mismo día y lote), por suma."""
    gb, gf = defaultdict(list), defaultdict(list)
    for b in banco:
        if b["match"] is None and b["origen"] == "Mes":
            gb[(b["fecha"], b["causal"], b["importe"] > 0)].append(b)
    for f in fbs:
        if f["match"] is None and f["origen"] == "Mes":
            gf[(f["fecha"], _clave_lote(f), f["importe"] > 0)].append(f)
            gf[(f["fecha"], "*", f["importe"] > 0)].append(f)
    for (fecha, _, signo), bs in sorted(gb.items(), key=lambda kv: kv[0][0]):
        if any(b["match"] is not None for b in bs):
            continue
        total = sum(b["importe"] for b in bs)
        for (ff, _, sf), fs in gf.items():
            if sf == signo and abs((ff - fecha).days) <= dias and all(f["match"] is None for f in fs) \
                    and abs(sum(f["importe"] for f in fs) - total) <= TOLERANCIA and (len(bs) > 1 or len(fs) > 1):
                _unir(bs, fs, "Lote (suma del día)")
                break


def _gastos(banco, fbs):
    """El mes de gastos/impuestos bancarios se registra en FBS en pocos asientos globales."""
    b = [x for x in banco if x["match"] is None and x["origen"] == "Mes"
         and (x["cat"] == "GASTOS BANCARIOS" or x["causal"] in ("1684", "1685", "1297", "4145"))]
    f = [x for x in fbs if x["match"] is None and re.search(r"GASTOS BANCARIOS \d{2}/\d{4}", x["texto"].upper())]
    if b and f:
        dif = round(sum(x["importe"] for x in b) - sum(x["importe"] for x in f), 2)
        _unir(b, f, f"Asiento gastos bancarios (dif. {dif:,.2f})")
        return dif
    return None


def conciliar(banco, fbs):
    mismo_id = lambda b, f: (b["cuit"] and b["cuit"] == f["cuit"]) or (b["dni"] and b["dni"] == f["dni"])
    por_ref = lambda b, f: len(b["ref"]) >= 5 and b["ref"] in (f["texto"] + " " + f["ref"]).replace(".", "")
    por_nombre = lambda b, f: bool(b["nombre"] & f["nombre"])
    _pasada(banco, fbs, mismo_id, "1. CUIT/DNI", ventana=VENTANA_ID)
    _pasada(banco, fbs, por_ref, "2. Referencia", ventana=60)
    _pasada(banco, fbs, por_nombre, "3. Nombre", ventana=15)
    _agrupados(banco, fbs)
    dif_gastos = _gastos(banco, fbs)
    _compensaciones_fbs(fbs, misma_clave=True)
    _pasada(banco, fbs, lambda b, f: True, "4. Fecha + importe", ventana=0, unico=True)
    _lotes(banco, fbs)
    _pasada(banco, fbs, lambda b, f: True, "5. Sugerido (solo importe)", ventana=VENTANA_SUGERIDO)
    _pasada(banco, fbs, lambda b, f: b["cat"].startswith(("TRANSFERENCIA ENTRE", "INVERSIONES")),
            "5. Sugerido (importe en el mes)", ventana=31, unico=True)
    _combinaciones(banco, fbs)
    _compensaciones_fbs(fbs, misma_clave=False)
    _reversiones_fbs(fbs)
    return dif_gastos


def sector(p):
    if p["lado"] == "BANCO":
        return "S1" if p["importe"] > 0 else "S2"
    return "S3" if p["importe"] > 0 else "S4"


# ---------------------------------------------------------------- salida

def generar(salida, info_banco, movs, info_e, info_o, banco, fbs, dif_gastos, corte, manual=None):
    wb = Workbook()
    wb.remove(wb.active)
    pend = defaultdict(list)
    for p in banco + fbs:
        if p["match"] is None:
            pend[sector(p)].append(p)
    tot = {s: sum(abs(p["importe"]) for p in pend[s]) for s, *_ in SECTORES}
    saldo_e, saldo_o = info_e["saldo_final"], info_o["saldo_final"]
    saldo_fbs = saldo_e + saldo_o
    saldo_banco = movs[-1]["saldo"]
    ajustado = saldo_fbs + tot["S1"] - tot["S2"] - tot["S3"] + tot["S4"]

    # --- Tablero: solo saldos (el detalle de cada cuadro está en su propia hoja)
    ws = wb.create_sheet("Tablero")
    bold = Font(bold=True)
    borde = Border(bottom=Side(style="thin", color="999999"))
    ws["A1"] = f"CONCILIACIÓN BANCO MACRO — {info_banco.get('Empresa', '')} — al {corte:%d/%m/%Y}"
    ws["A1"].font = Font(bold=True, size=14)
    man = manual or {}
    cab = ["Concepto", "Sistema"] + (["Conciliación manual", "Diferencia"] if man else [])
    for i, t in enumerate(cab, 1):
        c = ws.cell(3, i, t); c.font = Font(bold=True, color="FFFFFF"); c.fill = HEADER_FILL
    filas = [
        ("Cuenta E  " + info_e.get("cuenta", ""), saldo_e, man.get("E"), None),
        ("Cuenta O  " + info_o.get("cuenta", ""), saldo_o, man.get("O"), None),
        ("SALDO CONTABLE FBS (E + O)", saldo_fbs, man and man["E"] + man["O"], "total"),
        None,
    ]
    for s_, nombre, signo in SECTORES:
        filas.append((f"({signo}) {nombre}  — {len(pend[s_])} partidas, ver hoja {HOJAS[s_]}", tot[s_], man.get(s_), None))
    aj_man = man and man["E"] + man["O"] + man["S1"] - man["S2"] - man["S3"] + man["S4"]
    filas += [
        None,
        ("SALDO CONTABLE AJUSTADO", ajustado, aj_man, "total"),
        (f"SALDO SEGÚN EXTRACTO al {corte:%d/%m/%Y}", saldo_banco, man.get("banco"), "total"),
        ("DIFERENCIA (control)", ajustado - saldo_banco, man and aj_man - man["banco"], "dif"),
    ]
    r = 4
    for fila in filas:
        if fila:
            texto, valor, manual_v, estilo = fila
            ws.cell(r, 1, texto)
            vals = [valor] + ([manual_v, None if manual_v is None else valor - manual_v] if man else [])
            for i, v in enumerate(vals, 2):
                if v is not None:
                    c = ws.cell(r, i, round(v, 2)); c.number_format = NUM_FMT
            for i in range(1, len(cab) + 1):
                ws.cell(r, i).border = borde
                if estilo:
                    ws.cell(r, i).font = bold
            if estilo == "dif":
                fill = PatternFill("solid", fgColor="C6EFCE" if abs(valor) < 1000 else "FFC7CE")
                for i in range(1, len(cab) + 1):
                    ws.cell(r, i).fill = fill
        r += 1
    r += 1
    ws.cell(r, 1, "Calidad del cruce — movimientos del extracto del mes").font = bold
    ws.cell(r, 2, "Cant.").font = bold; ws.cell(r, 3, "Importe").font = bold
    r += 1
    metodos = defaultdict(lambda: [0, 0])
    for p in banco:
        if p["origen"] == "Mes":
            k = p["metodo"].split(" (dif")[0] if p["match"] is not None else "Sin conciliar"
            metodos[k][0] += 1; metodos[k][1] += abs(p["importe"])
    for k in sorted(metodos):
        ws.cell(r, 1, "  " + k); ws.cell(r, 2, metodos[k][0]); c = ws.cell(r, 3, metodos[k][1]); c.number_format = NUM_FMT
        r += 1
    if dif_gastos is not None:
        ws.cell(r, 1, "  Diferencia asiento gastos bancarios (banco − FBS)"); c = ws.cell(r, 3, dif_gastos)
        c.number_format = NUM_FMT
    for col, w in zip("ABCD", (78, 20, 20, 18)):
        ws.column_dimensions[col].width = w

    # --- Detalle de cada sector
    for s, nombre, _ in SECTORES:
        items = sorted(pend[s], key=lambda p: p["fecha"])
        hoja(wb, HOJAS[s], ["Fecha", "Días", "Origen", "Concepto / Comprobante", "Importe", "Categoría", "CUIT/DNI",
                     "Asiento FBS / Ref. banco"],
             [[p["fecha"], (corte - p["fecha"]).days, p["origen"], p["texto"], abs(p["importe"]), p["cat"],
               p["cuit"] or p["dni"], p.get("asiento") or p["ref"]] for p in items],
             [11, 7, 10, 60, 18, 32, 14, 18], {"A": "dd/mm/yyyy", "E": NUM_FMT})
        wd = wb[HOJAS[s]]
        wd.insert_rows(1)
        wd["A1"] = f"{nombre} — {len(items)} partidas — total {tot[s]:,.2f}"
        wd["A1"].font = bold
        wd.freeze_panes = "A3"
        wd.auto_filter.ref = f"A2:H{len(items) + 2}"

    # --- Conciliados
    filas, vistos = [], set()
    for b in banco:
        if b["match"] is None or id(b["match"]) in vistos:
            continue
        grupo_f = b["match"]
        vistos.add(id(grupo_f))
        grupo_b = grupo_f[0]["match"] if grupo_f else [b]
        n = max(len(grupo_b), len(grupo_f))
        for i in range(n):
            x = grupo_b[i] if i < len(grupo_b) else None
            y = grupo_f[i] if i < len(grupo_f) else None
            filas.append([b["metodo"],
                          x and x["fecha"], x and x["texto"], x and x["importe"], x and x["origen"],
                          y and y["fecha"], y and y["texto"], y and y["importe"], y and y.get("asiento"),
                          y and y["origen"]])
    for f in fbs:
        if f["match"] == []:
            filas.append([f["metodo"], None, None, None, None, f["fecha"], f["texto"], f["importe"], f.get("asiento"),
                          f["origen"]])
    hoja(wb, "Conciliados", ["Método", "Fecha banco", "Concepto banco", "Importe banco", "Origen banco",
                             "Fecha FBS", "Comprobante FBS", "Importe FBS", "Asiento FBS", "Origen FBS"],
         filas, [30, 11, 45, 16, 9, 11, 55, 16, 11, 9],
         {"B": "dd/mm/yyyy", "D": NUM_FMT, "F": "dd/mm/yyyy", "H": NUM_FMT})
    wb.save(salida)
    return tot, ajustado, saldo_banco


def main(extracto, mayor_e, mayor_o, salida, anterior=None):
    info_banco, movs = leer(extracto)
    info_e, asientos_e = leer_mayor(mayor_e)
    info_o, asientos_o = leer_mayor(mayor_o)
    desde, corte = movs[0]["fecha"].replace(day=1), movs[-1]["fecha"]
    banco = [partida_banco(m) for m in movs]
    fbs = [partida_fbs(a) for a in movimientos_fbs(asientos_e, asientos_o)]
    if anterior:
        for p in leer_arrastre(anterior, desde):
            (banco if p["lado"] == "BANCO" else fbs).append(p)
    dif_gastos = conciliar(banco, fbs)
    manual = leer_totales_manual(anterior) if anterior else None
    return generar(salida, info_banco, movs, info_e, info_o, banco, fbs, dif_gastos, corte, manual), banco, fbs


if __name__ == "__main__":
    (tot, ajustado, saldo_banco), _, _ = main(*sys.argv[1:])
    print({k: round(v, 2) for k, v in tot.items()}, "ajustado", round(ajustado, 2), "banco", saldo_banco)
